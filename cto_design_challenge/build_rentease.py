# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "RentEase_CTO_Strategy.pptx")

NAVY = RGBColor(0x0B, 0x24, 0x47)
GREEN = RGBColor(0x1B, 0x7F, 0x5B)
ORANGE = RGBColor(0xE0, 0x6A, 0x1F)
BLUE = RGBColor(0x46, 0x58, 0xC9)
PURPLE = RGBColor(0x5E, 0x41, 0x98)
GRAY = RGBColor(0x5B, 0x64, 0x72)
TEXT = RGBColor(0x22, 0x22, 0x2A)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF8)
BORDER = RGBColor(0xD7, 0xDC, 0xE8)

HEAD_FONT = "Century Gothic"
BODY_FONT = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_gradient_bar(slide, x, y, w, h, c1=RGBColor(0xFF, 0x65, 0x01), c2=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.line.fill.background()
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        fill.gradient_angle = 0.0
    except Exception:
        pass
    return shp


def add_rect(slide, x, y, w, h, color, line=False, round_=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.shadow.inherit = False
    if line:
        shp.line.color.rgb = BORDER
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp


def add_text(slide, x, y, w, h, text, size=14, color=TEXT, bold=False,
             italic=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=13, color=TEXT,
                 font=BODY_FONT, bullet_color=ORANGE, space_after=8,
                 line_spacing=1.05):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (label, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.size = Pt(size)
        r0.font.name = font
        r0.font.color.rgb = bullet_color
        if label:
            r1 = p.add_run()
            r1.text = label + "  "
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.italic = True
            r1.font.name = font
            r1.font.color.rgb = NAVY
        if desc:
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(size)
            r2.font.italic = True
            r2.font.name = font
            r2.font.color.rgb = color
    return box


def scaffold(slide, kicker, title, page_no, subtitle="CTO Case Study Response"):
    header_h = Inches(0.62)
    add_rect(slide, 0, 0, SW, header_h, BLACK)
    add_text(slide, Inches(0.4), 0, Inches(7), header_h,
              "RENTEASE — SCALING STRATEGY", size=13, color=WHITE,
              bold=True, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, SW - Inches(4.4), 0, Inches(4.0), header_h, subtitle,
              size=12, color=RGBColor(0xE8, 0xE8, 0xEC), italic=True,
              font=BODY_FONT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_gradient_bar(slide, 0, header_h, SW, Inches(0.06))
    add_gradient_bar(slide, 0, SH - Inches(0.09), SW, Inches(0.09))
    add_text(slide, SW - Inches(1.0), SH - Inches(0.42), Inches(0.7), Inches(0.3),
              str(page_no), size=11, color=NAVY, bold=True, align=PP_ALIGN.RIGHT)
    if kicker:
        add_text(slide, Inches(0.4), header_h + Inches(0.1), Inches(11.5), Inches(0.3),
                  kicker.upper(), size=11.5, color=ORANGE, bold=True, font=BODY_FONT)
        title_y = header_h + Inches(0.38)
    else:
        title_y = header_h + Inches(0.16)
    add_text(slide, Inches(0.4), title_y, Inches(12.5), Inches(0.55),
              title, size=23, color=NAVY, bold=True, font=HEAD_FONT)


def flow_row(slide, x, y, w, h, boxes, gap=0.2, box_color=NAVY,
             text_color=WHITE, arrow_color=GRAY, size=12):
    n = len(boxes)
    gap_in = Inches(gap)
    box_w = (w - gap_in * (n - 1)) / n
    cx = x
    for i, label in enumerate(boxes):
        color = box_color[i] if isinstance(box_color, list) else box_color
        add_rect(slide, cx, y, box_w, h, color, round_=True)
        add_text(slide, cx + Inches(0.06), y, box_w - Inches(0.12), h, label,
                  size=size, color=text_color, bold=True, font=BODY_FONT,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
        if i < n - 1:
            add_text(slide, cx + box_w, y, gap_in, h, "→", size=18,
                      color=arrow_color, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
        cx += box_w + gap_in
    return box_w


# ========================================================= SLIDE 1: TITLE ==
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
band_h = Inches(3.3)
add_rect(s, 0, 0, SW, band_h, BLACK)
add_gradient_bar(s, 0, band_h, SW, Inches(0.1))
add_gradient_bar(s, 0, SH - Inches(0.12), SW, Inches(0.12))
add_text(s, Inches(0.6), band_h / 2 - Inches(0.8), SW - Inches(1.2), Inches(1.0),
          "Scaling RentEase", size=44, color=WHITE, bold=True,
          font=HEAD_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.6), band_h / 2 + Inches(0.22), SW - Inches(1.2), Inches(0.6),
          "A Technology Strategy for Trusted, Capital-Efficient Growth", size=17,
          color=RGBColor(0xFF, 0x9A, 0x4D), italic=True, font=BODY_FONT,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

fy = band_h + Inches(0.9)
for label, val in [
    ("Perspective", "Chief Technology Officer"),
    ("Prepared By", "Keshav"),
    ("Case", "RentEase Leadership Case Study — Scaling RentEase"),
]:
    box = s.shapes.add_textbox(Inches(1.2), fy, Inches(11), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = label + " :  "
    r1.font.bold = True; r1.font.size = Pt(15); r1.font.name = BODY_FONT
    r1.font.color.rgb = NAVY
    r2 = p.add_run(); r2.text = val
    r2.font.italic = True; r2.font.size = Pt(15); r2.font.name = BODY_FONT
    r2.font.color.rgb = TEXT
    fy += Inches(0.55)

# ============================================ SLIDE 2: CTO LENS ============
s = prs.slides.add_slide(BLANK)
scaffold(s, "Framing", "The Two Challenges That Are Actually a Tech Problem", 2)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.6),
          "Six challenges are on the table. Only two are mine to fix directly — the rest are "
          "pricing, marketing, or capital-strategy calls that technology can support but not own.",
          size=13, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.15)

chal_cols = [
    (GREEN, "Damage, Loss &\nEarly Returns",
     "Every damaged, lost, or early-returned item is margin leaking out of the business with "
     "no system catching it early or assigning responsibility fairly."),
    (ORANGE, "Rising Delivery,\nMaintenance & Repair Cost",
     "Cost-per-order for logistics and upkeep is climbing with volume instead of falling — "
     "the sign of an operation running on manual coordination, not a system."),
]
cx = Inches(0.4)
cw = Inches(6.0)
for color, title, desc in chal_cols:
    add_rect(s, cx, Inches(2.3), cw, Inches(0.85), color, round_=True)
    add_text(s, cx + Inches(0.2), Inches(2.3), cw - Inches(0.4), Inches(0.85), title,
              size=16, color=WHITE, bold=True, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE,
              line_spacing=1.05)
    add_rect(s, cx, Inches(3.25), cw, Inches(1.5), LIGHT_BG)
    add_text(s, cx + Inches(0.2), Inches(3.4), cw - Inches(0.4), Inches(1.2), desc,
              size=12, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.2)
    cx += cw + Inches(0.5)

add_text(s, Inches(0.4), Inches(5.0), Inches(12.5), Inches(0.35),
          "Working Assumptions", size=15, color=PURPLE, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(5.45), Inches(12.5), Inches(1.6), [
    ("Current state:", "2 cities, 25,000+ users — inventory and delivery are likely coordinated "
     "manually or with basic tooling, not a purpose-built ops platform."),
    ("Scale target:", "multi-city national expansion, meaning every manual process today becomes "
     "a linear cost driver tomorrow unless it's systematized first."),
], size=12.5)

# =============================================== SLIDE 3: CORE RECOMMENDATION
s = prs.slides.add_slide(BLANK)
scaffold(s, "Strategic Recommendation", "Build “RentOS”: The Platform Before the Cities", 3)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.55),
          "One recommendation, not five: before spending on a third or fourth city, build the "
          "internal platform that makes every future city cheaper to run than the last one.",
          size=13, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.15)

layers = ["Asset Tags\n(QR / RFID)", "Condition Capture\n(photo + CV score)",
          "Inventory &\nRouting Engine", "Predictive\nMaintenance", "Expansion\nAnalytics"]
flow_row(s, Inches(0.4), Inches(2.3), Inches(12.5), Inches(1.0), layers,
          box_color=[NAVY, GREEN, ORANGE, ORANGE, PURPLE], size=11.5)

add_text(s, Inches(0.4), Inches(3.55), Inches(6.0), Inches(0.35),
          "Why a Platform, Not a Feature List", size=14, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(4.0), Inches(6.0), Inches(3.0), [
    ("Every layer is reusable", "across every new city — the marginal tech cost of city #5 is near zero once RentOS exists."),
    ("Each layer targets one challenge:", "tags + capture fix damage/loss; routing + maintenance fix rising cost per order."),
    ("It's sequenced, not simultaneous:", "tags and capture ship first because they generate the data the other layers depend on."),
], size=12)

add_rect(s, Inches(6.9), Inches(3.55), Inches(6.0), Inches(3.45), LIGHT_BG)
add_text(s, Inches(7.1), Inches(3.7), Inches(5.6), Inches(0.3),
          "The Pitch, in One Line", size=13.5, color=PURPLE, bold=True, font=HEAD_FONT)
add_text(s, Inches(7.1), Inches(4.1), Inches(5.6), Inches(2.7),
          "RentEase doesn't need to outspend Ikea or Amazon on warehouses to win rental — it "
          "needs to know, better than they ever will, exactly what condition every item is in "
          "and exactly when it will next fail. That knowledge is what RentOS is built to capture, "
          "and it compounds with every rental RentEase completes.",
          size=12.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.25)

# ========================================= SLIDE 4: DAMAGE / LOSS / RETURNS =
s = prs.slides.add_slide(BLANK)
scaffold(s, "Fix #1", "Closing the Damage, Loss & Early-Return Gap", 4)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
          "Today, a disputed damage claim is one person's word against another's. The fix is "
          "to make condition a timestamped fact, not a memory.",
          size=13, color=TEXT, italic=True, font=BODY_FONT)

flow4 = ["Scan-out at\npickup", "Photo + CV\ncondition score", "In customer's\npossession",
         "Photo + CV score\nat return", "Auto-compare &\nresolve deposit"]
flow_row(s, Inches(0.4), Inches(2.15), Inches(12.5), Inches(1.0), flow4,
          box_color=[NAVY, GREEN, GRAY, GREEN, ORANGE], size=11)

add_text(s, Inches(0.4), Inches(3.45), Inches(6.0), Inches(0.35),
          "What This Buys RentEase", size=14, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(3.9), Inches(6.0), Inches(3.0), [
    ("Fewer disputes,", "faster resolutions — condition delta is computed, not argued over."),
    ("Fair, automatic deposit handling,", "which is itself a trust signal to customers who are wary of hidden fees."),
    ("Loss reconciliation,", "asset tags mean a missing item is flagged at the next warehouse scan, not discovered months later."),
], size=12)

add_text(s, Inches(6.9), Inches(3.45), Inches(6.0), Inches(0.35),
          "Customer Trust, Not Just Cost Control", size=14, color=PURPLE, bold=True, font=HEAD_FONT)
add_rect(s, Inches(6.9), Inches(3.9), Inches(6.0), Inches(3.0), LIGHT_BG)
add_text(s, Inches(7.1), Inches(4.05), Inches(5.6), Inches(2.7),
          "The customers most resistant to renting over owning are worried about being blamed "
          "for pre-existing wear. A visible, photo-backed condition record at handover removes "
          "that fear before it's even raised — this is a retention lever disguised as a loss-"
          "prevention feature.",
          size=12.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.25)

# ============================================ SLIDE 5: COST AT SCALE =======
s = prs.slides.add_slide(BLANK)
scaffold(s, "Fix #2", "Bending the Delivery & Maintenance Cost Curve", 5)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
          "Cost per order should fall with scale, not rise with it — that only happens if "
          "routing and upkeep are systems, not judgment calls.",
          size=13, color=TEXT, italic=True, font=BODY_FONT)

add_text(s, Inches(0.4), Inches(2.25), Inches(6.0), Inches(0.35),
          "Two Systems to Put in Place", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(2.7), Inches(6.0), Inches(3.3), [
    ("Route-optimized batching:", "cluster pickups/deliveries by geography and time window instead of dispatching one order at a time."),
    ("Predictive maintenance:", "schedule servicing from usage hours/cycles per item, not a fixed calendar — catch failure before it becomes an emergency repair."),
    ("Partner-technician network app:", "an asset-light service layer — RentEase doesn't need to employ repair staff in every new city, just certify and route work to local partners through the app."),
], size=12)

add_text(s, Inches(6.9), Inches(2.25), Inches(6.0), Inches(0.35),
          "Why This Is Also the Expansion Answer", size=15, color=PURPLE, bold=True, font=HEAD_FONT)
add_rect(s, Inches(6.9), Inches(2.7), Inches(6.0), Inches(3.7), LIGHT_BG)
add_text(s, Inches(7.1), Inches(2.88), Inches(5.6), Inches(3.4),
          "The same partner-technician model that controls repair cost in existing cities is "
          "what makes new cities cheap to enter — RentEase brings the software and the "
          "certification standard, the local partner brings the labor and the vehicles. That's "
          "the difference between funding a fleet and funding an app, and it directly answers "
          "the \"expansion requires significant investment\" challenge without RentEase touching "
          "that line item.",
          size=12.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.25)

# ============================================ SLIDE 6: MOAT vs COMPETITORS =
s = prs.slides.add_slide(BLANK)
scaffold(s, "Creativity & Defensibility", "The Real Moat Isn't Inventory — It's Condition Data", 6)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.6),
          "Ikea and Amazon can out-spend RentEase on warehouses in a quarter. They cannot "
          "out-spend three years of item-level condition and failure data they don't have.",
          size=13, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.15)

mrows = [
    ("Large furniture / e-commerce entrants", GRAY, "Capital, brand reach, existing logistics networks."),
    ("RentOS condition + usage dataset", GREEN, "Knows exactly when an item will fail and what it's actually worth used — "
     "compounds with every rental and can't be bought or copied overnight."),
]
ry = Inches(2.35)
for label, color, desc in mrows:
    add_rect(s, Inches(0.4), ry, Inches(3.0), Inches(1.0), color, round_=True)
    add_text(s, Inches(0.55), ry, Inches(2.7), Inches(1.0), label, size=12.5, color=WHITE,
              bold=True, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    add_rect(s, Inches(3.6), ry, Inches(9.3), Inches(1.0), LIGHT_BG)
    add_text(s, Inches(3.8), ry, Inches(8.9), Inches(1.0), desc, size=12, color=TEXT,
              italic=True, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    ry += Inches(1.25)

add_text(s, Inches(0.4), Inches(5.0), Inches(12.5), Inches(0.35),
          "Where This Data Pays Off Beyond Cost Savings", size=14, color=PURPLE, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(5.45), Inches(12.5), Inches(1.5), [
    ("Dynamic, risk-based deposits:", "lower deposits for reliably-returned categories, building trust and conversion instead of friction."),
    ("A future B2B data product:", "resale-value and depreciation curves become licensable insight for insurers or manufacturers — optionality, not a Year-1 commitment."),
], size=12)

# ============================================ SLIDE 7: ROLLOUT & IMPACT ====
s = prs.slides.add_slide(BLANK)
scaffold(s, "Execution", "What Ships First, and the Number Behind It", 7)

phases = [
    ("Phase 1 — Now", GREEN, ["Asset tags + condition capture", "Live in existing 2 cities",
                               "Generates the dataset everything else depends on"]),
    ("Phase 2 — 6-12 mo.", ORANGE, ["Route optimization", "Predictive maintenance scheduling",
                                     "Cost-per-order starts declining, not rising"]),
    ("Phase 3 — New Cities", PURPLE, ["Partner-technician network app", "Asset-light entry per new city",
                                        "RentOS licensed to each new market, not rebuilt"]),
]
cw = Inches(3.95)
cx = Inches(0.4)
for label, color, items in phases:
    add_rect(s, cx, Inches(1.6), cw, Inches(0.55), color, round_=True)
    add_text(s, cx, Inches(1.6), cw, Inches(0.55), label, size=14.5, color=WHITE,
              bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, cx, Inches(2.25), cw, Inches(2.05), LIGHT_BG)
    iy = Inches(2.4)
    for it in items:
        add_text(s, cx + Inches(0.18), iy, cw - Inches(0.36), Inches(0.65),
                  "•  " + it, size=11, color=TEXT, italic=True, font=BODY_FONT,
                  line_spacing=1.1)
        iy += Inches(0.6)
    if cx + cw < SW - Inches(1):
        add_text(s, cx + cw, Inches(1.6), Inches(0.35), Inches(0.55), "→", size=18,
                  color=GRAY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += cw + Inches(0.35)

add_text(s, Inches(0.4), Inches(4.6), Inches(12.5), Inches(0.35),
          "Illustrative Impact (Assumptions Stated)", size=14, color=NAVY, bold=True, font=HEAD_FONT)
add_text(s, Inches(0.4), Inches(5.05), Inches(12.5), Inches(1.6),
          "Assume damage, loss, and early-return write-offs run at roughly 4-6% of rental revenue "
          "industry-wide — a plausible range for an unsystematized ops process at RentEase's stage. "
          "If condition capture and asset tagging recover even a third of that through fewer disputes, "
          "faster loss detection, and fairer deposit resolution, that is 1-2 points of revenue funding "
          "the next city's rollout before a single new rupee of investment is spent on it. The platform "
          "pays for its own expansion.",
          size=12.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.25)

prs.save(OUT)
print("Wrote", OUT)
