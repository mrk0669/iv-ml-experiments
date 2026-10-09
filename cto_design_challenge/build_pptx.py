# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image as PILImage

HERE = os.path.dirname(__file__)
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "Sustainable_Resource_Marketplace_Summary.pptx")

NAVY = RGBColor(0x0B, 0x24, 0x47)
GREEN = RGBColor(0x1B, 0x7F, 0x5B)
ORANGE = RGBColor(0xE0, 0x6A, 0x1F)
BLUE = RGBColor(0x46, 0x58, 0xC9)
PURPLE = RGBColor(0x5E, 0x41, 0x98)
GRAY = RGBColor(0x5B, 0x64, 0x72)
TEXT = RGBColor(0x22, 0x22, 0x2A)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF8)
BORDER = RGBColor(0xD7, 0xDC, 0xE8)

HEAD_FONT = "Century Gothic"
BODY_FONT = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_gradient_bar(slide, x, y, w, h, c1=RGBColor(0xFF, 0x65, 0x01), c2=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.line.fill.background()
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        fill.gradient_angle = 0.0
    except Exception:
        pass
    return shp


def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.shadow.inherit = False
    if line:
        shp.line.color.rgb = BORDER
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp


def add_text(slide, x, y, w, h, text, size=14, color=TEXT, bold=False,
             italic=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=13, color=TEXT,
                 font=BODY_FONT, bullet_color=ORANGE, space_after=8,
                 line_spacing=1.05):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (label, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.size = Pt(size)
        r0.font.name = font
        r0.font.color.rgb = bullet_color
        if label:
            r1 = p.add_run()
            r1.text = label + "  "
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.italic = True
            r1.font.name = font
            r1.font.color.rgb = NAVY
        if desc:
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(size)
            r2.font.italic = True
            r2.font.name = font
            r2.font.color.rgb = color
    return box


def add_picture_fit(slide, filename, x, y, max_w, max_h):
    path = os.path.join(ASSETS, filename)
    im = PILImage.open(path)
    iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    px = x + (max_w - w) // 2
    py = y + (max_h - h) // 2
    slide.shapes.add_picture(path, px, py, width=w, height=h)


def scaffold(slide, title, page_no, subtitle="Technical Design Report"):
    header_h = Inches(0.62)
    add_rect(slide, 0, 0, SW, header_h, BLACK)
    add_text(slide, Inches(0.4), 0, Inches(7), header_h,
              "SUSTAINABLE RESOURCE MARKETPLACE", size=13, color=WHITE,
              bold=True, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, SW - Inches(4.4), 0, Inches(4.0), header_h, subtitle,
              size=12, color=RGBColor(0xE8, 0xE8, 0xEC), italic=True,
              font=BODY_FONT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_gradient_bar(slide, 0, header_h, SW, Inches(0.06))
    add_gradient_bar(slide, 0, SH - Inches(0.09), SW, Inches(0.09))
    add_text(slide, SW - Inches(1.0), SH - Inches(0.42), Inches(0.7), Inches(0.3),
              str(page_no), size=11, color=NAVY, bold=True,
              align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.4), header_h + Inches(0.16), Inches(11.5), Inches(0.55),
              title, size=25, color=NAVY, bold=True, font=HEAD_FONT)


# ========================================================= SLIDE 1: TITLE ==
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
band_h = Inches(3.3)
add_rect(s, 0, 0, SW, band_h, BLACK)
add_gradient_bar(s, 0, band_h, SW, Inches(0.1))
add_gradient_bar(s, 0, SH - Inches(0.12), SW, Inches(0.12))
add_text(s, Inches(0.6), band_h / 2 - Inches(0.75), SW - Inches(1.2), Inches(1.0),
          "Sustainable Resource Marketplace", size=40, color=WHITE, bold=True,
          font=HEAD_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.6), band_h / 2 + Inches(0.28), SW - Inches(1.2), Inches(0.6),
          "A Technical Design Report — Condensed Summary", size=18,
          color=RGBColor(0xFF, 0x9A, 0x4D), italic=True, font=BODY_FONT,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

fy = band_h + Inches(0.9)
for label, val in [
    ("Report Scope", "User Experience & IA  |  System Architecture  |  Technology Stack & Data"),
    ("Prepared By", "Keshav"),
    ("Submission", "CTO Design Challenge — July 2026"),
]:
    box = s.shapes.add_textbox(Inches(1.2), fy, Inches(11), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = label + " :  "
    r1.font.bold = True; r1.font.size = Pt(15); r1.font.name = BODY_FONT
    r1.font.color.rgb = NAVY
    r2 = p.add_run(); r2.text = val
    r2.font.italic = True; r2.font.size = Pt(15); r2.font.name = BODY_FONT
    r2.font.color.rgb = TEXT
    fy += Inches(0.55)

# ============================================ SLIDE 2: EXEC SUMMARY / WHO ==
s = prs.slides.add_slide(BLANK)
scaffold(s, "Executive Summary & Who Uses the Platform", 2)
top = Inches(1.55)
add_text(s, Inches(0.4), top, Inches(12.5), Inches(0.9),
          "A B2B platform where businesses list, discover, and exchange surplus or "
          "reusable materials — turning waste liabilities into revenue while making "
          "the environmental impact of every transaction visible and auditable.",
          size=14, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.15)

add_text(s, Inches(0.4), Inches(2.55), Inches(4.0), Inches(0.4),
          "Guiding Principles", size=16, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(3.0), Inches(5.9), Inches(4.0), [
    ("Trust is the product.", "Verification, escrow, transparent metrics are core, not add-ons."),
    ("Sustainability data is first-class.", "CO2e and waste-diverted computed on every listing."),
    ("Modular before distributed.", "One deployable service, clear internal boundaries."),
    ("Buy commodity, build differentiation.", "Payments/KYC bought; matching & trust mechanics built."),
], size=12.5)

add_text(s, Inches(6.6), Inches(2.55), Inches(6.3), Inches(0.4),
          "Primary Personas", size=16, color=GREEN, bold=True, font=HEAD_FONT)
persona_rows = [
    ("Supplier (Seller)", "Turns surplus/waste into revenue or documented impact."),
    ("Buyer (Procurer)", "Sources verified, sustainable inputs and reports the impact."),
    ("Platform Admin", "Keeps the marketplace safe: verification, moderation, disputes."),
]
py = Inches(3.0)
for name, desc in persona_rows:
    add_rect(s, Inches(6.6), py, Inches(6.3), Inches(1.0), LIGHT_BG)
    add_text(s, Inches(6.8), py + Inches(0.08), Inches(5.9), Inches(0.35),
              name, size=13, color=NAVY, bold=True, font=BODY_FONT)
    add_text(s, Inches(6.8), py + Inches(0.44), Inches(5.9), Inches(0.5),
              desc, size=11.5, color=TEXT, italic=True, font=BODY_FONT)
    py += Inches(1.15)

# ===================================================== SLIDE 3: UX & IA ===
s = prs.slides.add_slide(BLANK)
scaffold(s, "User Experience & Information Architecture", 3)
add_picture_fit(s, "diagram_sitemap.png", Inches(0.3), Inches(1.5),
                 Inches(8.6), Inches(4.9))
add_text(s, Inches(9.1), Inches(1.55), Inches(3.9), Inches(0.35),
          "Core User Journey", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(9.1), Inches(2.0), Inches(3.9), Inches(2.3), [
    ("", "Discover -> Browse & filter -> View listing"),
    ("", "Send RFQ -> Negotiate terms"),
    ("", "Confirm & pay into escrow"),
    ("", "Shared fulfillment -> logistics -> receipt"),
    ("", "Rate, review & impact report generated"),
], size=11.5)
add_text(s, Inches(9.1), Inches(4.5), Inches(3.9), Inches(0.35),
          "UX Principles", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(9.1), Inches(4.9), Inches(3.9), Inches(1.7), [
    ("", "Impact metrics shown up front, not buried"),
    ("", "Low-friction, pre-tagged listing wizard"),
    ("", "Trust signals on every card and profile"),
], size=11.5)

# ================================================ SLIDE 4: ARCHITECTURE ===
s = prs.slides.add_slide(BLANK)
scaffold(s, "System Architecture", 4)
add_picture_fit(s, "diagram_architecture.png", Inches(0.3), Inches(1.5),
                 Inches(12.7), Inches(5.6))

# ============================================ SLIDE 5: STACK & DATA =======
s = prs.slides.add_slide(BLANK)
scaffold(s, "Technology Stack & Data Model", 5)
add_text(s, Inches(0.4), Inches(1.5), Inches(6.0), Inches(0.35),
          "Recommended Stack", size=15, color=GREEN, bold=True, font=HEAD_FONT)
stack_rows = [
    ("Frontend", "Next.js + TypeScript, Tailwind"),
    ("Backend", "Node.js (NestJS); Python/FastAPI for metrics"),
    ("Realtime", "WebSocket (Socket.IO)"),
    ("Event bus", "Amazon EventBridge + SQS"),
    ("Database", "PostgreSQL (RDS) + Redis + OpenSearch"),
    ("Storage / Payments", "S3 + CloudFront / Stripe Connect"),
    ("Hosting", "AWS ECS on Fargate, Terraform, GitHub Actions"),
]
ty = Inches(1.95)
for layer, tech in stack_rows:
    add_text(s, Inches(0.4), ty, Inches(1.7), Inches(0.4), layer, size=11.5,
              color=NAVY, bold=True, font=BODY_FONT)
    add_text(s, Inches(2.15), ty, Inches(4.3), Inches(0.4), tech, size=11.5,
              color=TEXT, italic=True, font=BODY_FONT)
    ty += Inches(0.55)

add_picture_fit(s, "diagram_datamodel.png", Inches(6.7), Inches(1.5),
                 Inches(6.3), Inches(5.6))

# ============================================ SLIDE 6: ROADMAP / CLOSE ====
s = prs.slides.add_slide(BLANK)
scaffold(s, "Roadmap & Key Trade-offs", 6)
add_text(s, Inches(0.4), Inches(1.55), Inches(6.0), Inches(0.35),
          "Phased Roadmap", size=15, color=GREEN, bold=True, font=HEAD_FONT)
roadmap = [
    ("MVP", "Web app, core marketplace, Stripe payments, manual logistics."),
    ("V2", "Mobile app, automated carrier APIs, recommendation engine."),
    ("V3", "Enterprise ERP integrations, marketplace-wide sustainability intel."),
]
ry = Inches(2.0)
for phase, desc in roadmap:
    add_rect(s, Inches(0.4), ry, Inches(6.0), Inches(1.15), LIGHT_BG)
    add_text(s, Inches(0.6), ry + Inches(0.1), Inches(1.2), Inches(0.35),
              phase, size=13, color=ORANGE, bold=True, font=HEAD_FONT)
    add_text(s, Inches(0.6), ry + Inches(0.48), Inches(5.6), Inches(0.6),
              desc, size=11.5, color=TEXT, italic=True, font=BODY_FONT)
    ry += Inches(1.3)

add_text(s, Inches(6.9), Inches(1.55), Inches(6.0), Inches(0.35),
          "Key Trade-offs & Assumptions", size=15, color=GREEN, bold=True,
          font=HEAD_FONT)
add_bullets(s, Inches(6.9), Inches(2.0), Inches(6.0), Inches(4.5), [
    ("Modular monolith,", "not microservices at launch, for engineering velocity."),
    ("Logistics is orchestration,", "not owned trucks/warehouses."),
    ("Category-level conversion factors,", "not per-item lifecycle assessment."),
    ("Monetization model,", "commission/subscription compatible with Stripe Connect, "
     "left as a product decision."),
], size=12)

prs.save(OUT)
print("Wrote", OUT)
