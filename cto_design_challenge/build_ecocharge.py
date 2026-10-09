# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image as PILImage

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "EcoCharge_CTO_Strategy.pptx")

NAVY = RGBColor(0x0B, 0x24, 0x47)
GREEN = RGBColor(0x1B, 0x7F, 0x5B)
ORANGE = RGBColor(0xE0, 0x6A, 0x1F)
BLUE = RGBColor(0x46, 0x58, 0xC9)
PURPLE = RGBColor(0x5E, 0x41, 0x98)
GRAY = RGBColor(0x5B, 0x64, 0x72)
TEXT = RGBColor(0x22, 0x22, 0x2A)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF4, 0xF8)
BORDER = RGBColor(0xD7, 0xDC, 0xE8)

HEAD_FONT = "Century Gothic"
BODY_FONT = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_gradient_bar(slide, x, y, w, h, c1=RGBColor(0xFF, 0x65, 0x01), c2=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.line.fill.background()
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        fill.gradient_angle = 0.0
    except Exception:
        pass
    return shp


def add_rect(slide, x, y, w, h, color, line=False, round_=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.shadow.inherit = False
    if line:
        shp.line.color.rgb = BORDER
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    return shp


def add_text(slide, x, y, w, h, text, size=14, color=TEXT, bold=False,
             italic=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=13, color=TEXT,
                 font=BODY_FONT, bullet_color=ORANGE, space_after=8,
                 line_spacing=1.05):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (label, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r0 = p.add_run()
        r0.text = "●  "
        r0.font.size = Pt(size)
        r0.font.name = font
        r0.font.color.rgb = bullet_color
        if label:
            r1 = p.add_run()
            r1.text = label + "  "
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.italic = True
            r1.font.name = font
            r1.font.color.rgb = NAVY
        if desc:
            r2 = p.add_run()
            r2.text = desc
            r2.font.size = Pt(size)
            r2.font.italic = True
            r2.font.name = font
            r2.font.color.rgb = color
    return box


def scaffold(slide, kicker, title, page_no, subtitle="CTO Case Study Response"):
    header_h = Inches(0.62)
    add_rect(slide, 0, 0, SW, header_h, BLACK)
    add_text(slide, Inches(0.4), 0, Inches(7), header_h,
              "ECOCHARGE — SCALING STRATEGY", size=13, color=WHITE,
              bold=True, font=BODY_FONT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, SW - Inches(4.4), 0, Inches(4.0), header_h, subtitle,
              size=12, color=RGBColor(0xE8, 0xE8, 0xEC), italic=True,
              font=BODY_FONT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_gradient_bar(slide, 0, header_h, SW, Inches(0.06))
    add_gradient_bar(slide, 0, SH - Inches(0.09), SW, Inches(0.09))
    add_text(slide, SW - Inches(1.0), SH - Inches(0.42), Inches(0.7), Inches(0.3),
              str(page_no), size=11, color=NAVY, bold=True, align=PP_ALIGN.RIGHT)
    if kicker:
        add_text(slide, Inches(0.4), header_h + Inches(0.1), Inches(11.5), Inches(0.3),
                  kicker.upper(), size=11.5, color=ORANGE, bold=True, font=BODY_FONT)
        title_y = header_h + Inches(0.38)
    else:
        title_y = header_h + Inches(0.16)
    add_text(slide, Inches(0.4), title_y, Inches(12.5), Inches(0.55),
              title, size=24, color=NAVY, bold=True, font=HEAD_FONT)


def flow_row(slide, x, y, w, h, boxes, gap=0.25, box_color=NAVY,
             text_color=WHITE, arrow_color=GRAY, size=12):
    n = len(boxes)
    gap_in = Inches(gap)
    box_w = (w - gap_in * (n - 1)) / n
    cx = x
    for i, label in enumerate(boxes):
        color = box_color[i] if isinstance(box_color, list) else box_color
        add_rect(slide, cx, y, box_w, h, color, round_=True)
        add_text(slide, cx + Inches(0.08), y, box_w - Inches(0.16), h, label,
                  size=size, color=text_color, bold=True, font=BODY_FONT,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
        if i < n - 1:
            add_text(slide, cx + box_w, y, gap_in, h, "→", size=20,
                      color=arrow_color, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
        cx += box_w + gap_in
    return box_w


# ========================================================= SLIDE 1: TITLE ==
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, WHITE)
band_h = Inches(3.3)
add_rect(s, 0, 0, SW, band_h, BLACK)
add_gradient_bar(s, 0, band_h, SW, Inches(0.1))
add_gradient_bar(s, 0, SH - Inches(0.12), SW, Inches(0.12))
add_text(s, Inches(0.6), band_h / 2 - Inches(0.8), SW - Inches(1.2), Inches(1.0),
          "Scaling EcoCharge", size=42, color=WHITE, bold=True,
          font=HEAD_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(0.6), band_h / 2 + Inches(0.22), SW - Inches(1.2), Inches(0.6),
          "A Technology Roadmap for National Scale-Up", size=18,
          color=RGBColor(0xFF, 0x9A, 0x4D), italic=True, font=BODY_FONT,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

fy = band_h + Inches(0.9)
for label, val in [
    ("Perspective", "Chief Technology Officer"),
    ("Prepared By", "Keshav"),
    ("Case", "EcoCharge Leadership Case Study — Scaling EcoCharge"),
]:
    box = s.shapes.add_textbox(Inches(1.2), fy, Inches(11), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = label + " :  "
    r1.font.bold = True; r1.font.size = Pt(15); r1.font.name = BODY_FONT
    r1.font.color.rgb = NAVY
    r2 = p.add_run(); r2.text = val
    r2.font.italic = True; r2.font.size = Pt(15); r2.font.name = BODY_FONT
    r2.font.color.rgb = TEXT
    fy += Inches(0.55)

# ================================================ SLIDE 2: CTO MANDATE ====
s = prs.slides.add_slide(BLANK)
scaffold(s, "Framing", "The CTO's Mandate in This Scale-Up", 2)
add_text(s, Inches(0.4), Inches(1.55), Inches(12.5), Inches(0.75),
          "My job is not to fix marketing, financing, or reverse logistics — it's to make sure the "
          "technology can be trusted, manufactured, and defended at 10x today's volume. Every "
          "recommendation below is filtered through that lens.",
          size=13.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.15)

add_text(s, Inches(0.4), Inches(2.5), Inches(6.0), Inches(0.4),
          "Working Assumptions", size=16, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(2.95), Inches(6.0), Inches(3.6), [
    ("Current stage:", "lab-validated, pilot-deployed with institutional customers (stationary storage, not yet consumer)."),
    ("Team:", "4 engineering co-founders — no dedicated manufacturing or quality function yet."),
    ("Feedstock:", "industrial + e-waste inputs, meaning material composition varies by batch and by region."),
    ("Capital:", "early growth-stage funding — enough for one major scale-up bet, not several in parallel."),
], size=12.5)

add_text(s, Inches(6.9), Inches(2.5), Inches(6.0), Inches(0.4),
          "Guiding Principles", size=16, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(6.9), Inches(2.95), Inches(6.0), Inches(3.6), [
    ("Reliability before reach.", "No mass-market push until failure and lifespan data exist to back it."),
    ("Design for manufacturability,", "not just for the lab bench, from the next revision onward."),
    ("Vertical-integrate only the moat:", "core chemistry & QC in-house, commodity steps out."),
    ("Protect what's actually defensible,", "not everything — patents are expensive and slow."),
], size=12.5)

# ============================================== SLIDE 3: TRL / RELIABILITY =
s = prs.slides.add_slide(BLANK)
scaffold(s, "Technical Challenges", "Closing the Reliability Gap Before Mass Production", 3)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.55),
          "Customer skepticism about lifespan is a data problem, not a messaging problem — "
          "marketing can't sell what engineering hasn't proven yet.",
          size=13, color=TEXT, italic=True, font=BODY_FONT)

add_text(s, Inches(0.4), Inches(2.15), Inches(6.0), Inches(0.35),
          "Technology Readiness Ladder", size=15, color=GREEN, bold=True, font=HEAD_FONT)
trl_labels = ["TRL 6\nPilot\nvalidated", "TRL 7\nField data\n(12+ mo)", "TRL 8\nCert. &\nprocess lock",
              "TRL 9\nMass mfg.\nproven"]
trl_colors = [GREEN, ORANGE, NAVY, NAVY]
flow_row(s, Inches(0.4), Inches(2.6), Inches(8.6), Inches(1.1), trl_labels,
          box_color=trl_colors, size=12)
add_text(s, Inches(0.4), Inches(3.9), Inches(8.6), Inches(0.3),
          "We are here — the gap to TRL 9 is what this roadmap closes over 3 years.",
          size=11, color=GRAY, italic=True, font=BODY_FONT)

add_text(s, Inches(0.4), Inches(4.4), Inches(8.6), Inches(0.35),
          "Priority Work Before Scale-Up", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(4.85), Inches(8.6), Inches(2.2), [
    ("Accelerated + real-world aging tests", "to publish an honest cycle-life and lifespan number, not a lab-only figure."),
    ("Third-party safety certification", "(UL 1973 / IEC 62619-class testing) to give customers and insurers an external trust signal."),
    ("Feedstock-variability testing", "so performance is proven across the realistic range of waste-material inputs, not one clean batch."),
], size=12)

add_text(s, Inches(9.3), Inches(2.15), Inches(3.6), Inches(0.35),
          "Why This Is the CTO's Fix", size=14, color=PURPLE, bold=True, font=HEAD_FONT)
add_rect(s, Inches(9.3), Inches(2.6), Inches(3.6), Inches(4.0), LIGHT_BG)
add_text(s, Inches(9.5), Inches(2.78), Inches(3.2), Inches(3.7),
          "Every trust claim CMO makes about \"reliable batteries from waste\" has to be "
          "backed by a test report. That report is engineering's deliverable, not "
          "marketing's — so reliability data is treated here as a Year-1 R&D output, "
          "not an afterthought.",
          size=12, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.2)

# ============================================ SLIDE 4: MAKE VS BUY =========
s = prs.slides.add_slide(BLANK)
scaffold(s, "Manufacturing Strategy", "Build, Outsource, or Partner — Not All Three Everywhere", 4)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.55),
          "With limited capital, full in-house manufacturing is the wrong first bet. The rule: "
          "keep the parts that are the moat, buy the parts that are commodity.",
          size=13, color=TEXT, italic=True, font=BODY_FONT)

cols = [
    ("Keep In-House", GREEN, "Core material process\n& formulation",
     ["Waste sorting / purification recipe", "Electrode formulation & mixing", "Final QC & cell testing"],
     "This is the patent. Outsourcing it hands a competitor the moat."),
    ("Outsource", ORANGE, "Commodity manufacturing\nsteps",
     ["Cell casing / packaging", "PCB & BMS electronics assembly", "Pack-level assembly"],
     "Contract manufacturers already run these at scale, cheaper and faster than we can build."),
    ("Partner", PURPLE, "Capital equipment &\nscale-up",
     ["Pilot-to-mass-production line design", "Regional assembly partners as volume grows", "Equipment financing / leasing"],
     "Avoids a large upfront capex bet before demand is proven regionally."),
]
cw = Inches(3.9)
cx = Inches(0.4)
for title, color, sub, items, note in cols:
    add_rect(s, cx, Inches(2.25), cw, Inches(0.75), color, round_=True)
    add_text(s, cx + Inches(0.15), Inches(2.3), cw - Inches(0.3), Inches(0.65), title,
              size=16, color=WHITE, bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, cx + Inches(0.15), Inches(3.08), cw - Inches(0.3), Inches(0.5), sub,
              size=11.5, color=NAVY, bold=True, italic=True, font=BODY_FONT,
              align=PP_ALIGN.CENTER, line_spacing=1.05)
    add_rect(s, cx, Inches(3.65), cw, Inches(2.2), LIGHT_BG)
    iy = Inches(3.78)
    for it in items:
        add_text(s, cx + Inches(0.18), iy, cw - Inches(0.36), Inches(0.5),
                  "•  " + it, size=11, color=TEXT, font=BODY_FONT, line_spacing=1.05)
        iy += Inches(0.55)
    add_text(s, cx + Inches(0.15), Inches(6.0), cw - Inches(0.3), Inches(0.9), note,
              size=10.5, color=GRAY, italic=True, font=BODY_FONT, line_spacing=1.1)
    cx += cw + Inches(0.25)

# ============================================ SLIDE 5: QUALITY AT SCALE ===
s = prs.slides.add_slide(BLANK)
scaffold(s, "Operations", "Maintaining Quality as Production Scales", 5)
add_text(s, Inches(0.4), Inches(1.5), Inches(12.5), Inches(0.5),
          "Quality can't be inspected in at the end of the line — it has to be a closed loop "
          "that feeds back into R&D.", size=13, color=TEXT, italic=True, font=BODY_FONT)

loop_labels = ["Incoming Material QC\n(waste feedstock)", "In-Process SPC\n(statistical control)",
               "Outgoing & Field QC\n(cert. + telemetry)", "Feedback to R&D\n(design update)"]
flow_row(s, Inches(0.4), Inches(2.3), Inches(12.5), Inches(1.1), loop_labels,
          box_color=[GREEN, NAVY, ORANGE, PURPLE], size=13)
add_text(s, Inches(0.4), Inches(3.55), Inches(12.5), Inches(0.3),
          "The loop closes: field QC and BMS telemetry data flow back into the next material/design revision.",
          size=11, color=GRAY, italic=True, font=BODY_FONT)

add_text(s, Inches(0.4), Inches(4.15), Inches(6.0), Inches(0.35),
          "Systems to Put in Place", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(4.6), Inches(6.0), Inches(2.4), [
    ("Digital batch traceability", "— every cell traceable to its feedstock batch and process run, essential given input variability."),
    ("Statistical process control (SPC)", "on the 2-3 process steps most sensitive to feedstock variation."),
    ("Certification roadmap", "aligned to target markets (safety + performance standards) before scale-up, not after."),
], size=12)

add_text(s, Inches(6.9), Inches(4.15), Inches(6.0), Inches(0.35),
          "Why Traceability Is Non-Negotiable Here", size=15, color=PURPLE, bold=True, font=HEAD_FONT)
add_rect(s, Inches(6.9), Inches(4.6), Inches(6.0), Inches(2.4), LIGHT_BG)
add_text(s, Inches(7.1), Inches(4.78), Inches(5.6), Inches(2.1),
          "Unlike a conventional battery maker with uniform raw material, our feedstock is "
          "heterogeneous by design. Batch-level traceability is what lets us isolate a "
          "quality issue to a specific waste-material source instead of scrapping entire "
          "production runs — it's a cost-control tool as much as a quality one.",
          size=12, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.2)

# ============================================ SLIDE 6: R&D + IP ===========
s = prs.slides.add_slide(BLANK)
scaffold(s, "R&D & IP Strategy", "Where to Invest Engineering Effort — and What to Protect", 6)

add_text(s, Inches(0.4), Inches(1.5), Inches(6.0), Inches(0.4),
          "R&D Priorities (in order)", size=15, color=GREEN, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.4), Inches(1.95), Inches(6.0), Inches(4.5), [
    ("1. Feedstock flexibility.", "Widen the range of acceptable waste inputs — turns raw-material availability from a risk into a moat."),
    ("2. Performance parity.", "Close any remaining energy-density gap vs. incumbent Li-ion for the target commercial segments."),
    ("3. BMS & monitoring software.", "A software layer for safety and performance is a harder-to-copy, faster-to-iterate moat than the cell alone."),
    ("4. Design for second life.", "Build recyclability into the next design revision so reverse logistics (Ops' problem) is cheaper to run."),
], size=12.5)

add_text(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(0.4),
          "IP Protection Strategy", size=15, color=GREEN, bold=True, font=HEAD_FONT)
ip_rows = [
    ("Patent", "Core material formulation & waste-purification process (already filed) — the actual moat."),
    ("Patent (file next)", "BMS/software algorithms and any cell-chemistry variants developed during scale-up."),
    ("Trade secret", "Sourcing relationships & feedstock-sorting know-how — faster to protect, hard to patent usefully."),
    ("Defensive publish", "Peripheral process tweaks not worth patenting, to block competitors from patenting them first."),
]
ry = Inches(1.95)
for label, desc in ip_rows:
    add_rect(s, Inches(6.9), ry, Inches(6.0), Inches(1.05), LIGHT_BG)
    add_text(s, Inches(7.1), ry + Inches(0.08), Inches(5.6), Inches(0.3), label,
              size=12.5, color=ORANGE, bold=True, font=BODY_FONT)
    add_text(s, Inches(7.1), ry + Inches(0.4), Inches(5.6), Inches(0.6), desc,
              size=11, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.1)
    ry += Inches(1.15)

# ============================================ SLIDE 7: ROADMAP =============
s = prs.slides.add_slide(BLANK)
scaffold(s, "Summary", "3-Year Technology Roadmap", 7)

years = [
    ("Year 1", GREEN, ["Field-validate lifespan (12+ mo. data)", "Lock core process; outsource commodity mfg.", "File BMS/software patents"]),
    ("Year 2", ORANGE, ["Safety certification complete", "Scale via contract manufacturing partner", "Digital batch traceability live"]),
    ("Year 3", NAVY, ["Feedstock flexibility expanded", "Selective vertical integration where volume justifies capex", "Design-for-second-life ships"]),
]
cw = Inches(3.95)
cx = Inches(0.4)
for label, color, items in years:
    add_rect(s, cx, Inches(1.65), cw, Inches(0.6), color, round_=True)
    add_text(s, cx, Inches(1.65), cw, Inches(0.6), label, size=17, color=WHITE,
              bold=True, font=HEAD_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, cx, Inches(2.35), cw, Inches(2.5), LIGHT_BG)
    iy = Inches(2.5)
    for it in items:
        add_text(s, cx + Inches(0.2), iy, cw - Inches(0.4), Inches(0.7),
                  "•  " + it, size=11.5, color=TEXT, italic=True, font=BODY_FONT,
                  line_spacing=1.15)
        iy += Inches(0.78)
    if cx + cw < SW - Inches(1):
        add_text(s, cx + cw, Inches(1.65), Inches(0.35), Inches(0.6), "→", size=20,
                  color=GRAY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += cw + Inches(0.35)

add_text(s, Inches(0.4), Inches(5.15), Inches(12.5), Inches(0.35),
          "What Technology Contributes to Company Success", size=15, color=PURPLE,
          bold=True, font=HEAD_FONT)
add_text(s, Inches(0.4), Inches(5.6), Inches(12.5), Inches(1.2),
          "Reliability data gives Marketing a trust claim it can actually defend. A make/outsource "
          "split protects the moat while keeping capex within Finance's constraints. Feedstock "
          "flexibility turns the biggest supply risk into a durable advantage. That is the "
          "Technology function's contribution to scaling EcoCharge nationally.",
          size=12.5, color=TEXT, italic=True, font=BODY_FONT, line_spacing=1.2)

prs.save(OUT)
print("Wrote", OUT)
